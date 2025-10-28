"""
Document Regeneration Cloud Function
"""

from firebase_functions import https_fn
from firebase_admin import firestore, storage
from docx import Document
from datetime import datetime
import tempfile
import os

from .placeholders import replace_placeholders


@https_fn.on_call()
def regenerate_document(req: https_fn.CallableRequest) -> dict:
    """
    Regenerate an existing document.

    Uses the original generation_data to recreate the document.

    Request data:
        project_id: str
        document_id: str

    Returns:
        dict with success status and new document info
    """

    if not req.auth:
        raise https_fn.HttpsError(
            code=https_fn.FunctionsErrorCode.UNAUTHENTICATED,
            message='Authentication required'
        )

    project_id = req.data.get('project_id')
    document_id = req.data.get('document_id')

    if not project_id or not document_id:
        raise https_fn.HttpsError(
            code=https_fn.FunctionsErrorCode.INVALID_ARGUMENT,
            message='project_id and document_id are required'
        )

    db = firestore.client()
    bucket = storage.bucket()

    try:
        # Get project
        project_ref = db.collection('projects').document(project_id)
        project_doc = project_ref.get()

        if not project_doc.exists:
            raise https_fn.HttpsError(
                code=https_fn.FunctionsErrorCode.NOT_FOUND,
                message='Project not found'
            )

        project_data = project_doc.to_dict()

        # Find the document in generated_docs
        generated_docs = project_data.get('generated_docs', [])
        original_doc = None

        for doc in generated_docs:
            if doc['id'] == document_id:
                original_doc = doc
                break

        if not original_doc:
            raise https_fn.HttpsError(
                code=https_fn.FunctionsErrorCode.NOT_FOUND,
                message='Document not found'
            )

        # Get template
        template_id = original_doc['template_id']
        template_ref = db.collection('templates').document(template_id)
        template_doc = template_ref.get()

        if not template_doc.exists:
            raise https_fn.HttpsError(
                code=https_fn.FunctionsErrorCode.NOT_FOUND,
                message='Template not found'
            )

        template_data = template_doc.to_dict()

        # Use original generation data
        generation_data = original_doc.get('generation_data', {})

        # Download template
        template_path = template_data['file_path']
        template_blob = bucket.blob(template_path)

        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_template:
            template_blob.download_to_filename(temp_template.name)
            template_file = temp_template.name

        try:
            # Generate document
            doc = Document(template_file)
            replace_placeholders(doc, generation_data)

            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_output:
                doc.save(temp_output.name)
                output_file = temp_output.name

            try:
                # Delete old file
                old_path = original_doc.get('file_path')
                if old_path:
                    old_blob = bucket.blob(old_path)
                    try:
                        old_blob.delete()
                    except Exception as e:
                        print(f"Warning: Could not delete old file: {e}")

                # Upload new file
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_path = f"documents/{project_id}/{template_id}_{timestamp}.docx"
                output_blob = bucket.blob(output_path)
                output_blob.upload_from_filename(output_file)

                file_url = f"gs://{bucket.name}/{output_path}"
                file_size = os.path.getsize(output_file)

                # Update document info
                updated_doc = {
                    **original_doc,
                    'file_url': file_url,
                    'file_path': output_path,
                    'file_size': file_size,
                    'regenerated_at': firestore.SERVER_TIMESTAMP,
                    'regenerated_by': req.auth.uid
                }

                # Update in project
                # Remove old, add new
                new_generated_docs = [
                    doc if doc['id'] != document_id else updated_doc
                    for doc in generated_docs
                ]

                project_ref.update({
                    'generated_docs': new_generated_docs,
                    'updated_at': firestore.SERVER_TIMESTAMP
                })

                # Log activity
                db.collection('activities').add({
                    'action': 'regenerate_document',
                    'user_id': req.auth.uid,
                    'user_name': req.auth.token.get('name', 'Unknown'),
                    'resource_type': 'document',
                    'resource_id': document_id,
                    'resource_name': original_doc.get('template_name', ''),
                    'details': {
                        'project_id': project_id,
                        'template_id': template_id
                    },
                    'timestamp': firestore.SERVER_TIMESTAMP
                })

                return {
                    'success': True,
                    'document': updated_doc,
                    'download_url': file_url
                }

            finally:
                os.unlink(output_file)

        finally:
            os.unlink(template_file)

    except https_fn.HttpsError:
        raise
    except Exception as e:
        print(f"Error regenerating document: {e}")
        raise https_fn.HttpsError(
            code=https_fn.FunctionsErrorCode.INTERNAL,
            message=f'Internal error: {str(e)}'
        )

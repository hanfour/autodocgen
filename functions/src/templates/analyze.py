"""
Template Analysis Cloud Function
"""

from firebase_functions import https_fn
from firebase_admin import storage
from docx import Document
import tempfile

from ..documents.placeholders import find_placeholders


# Standard variables that are always available
STANDARD_VARIABLES = {
    'project_name', 'company_name', 'contact_name',
    'price', 'price_before_tax', 'tax_amount',
    'date', 'year', 'month', 'day',
    'roc_year', 'roc_date',
    'document_number', 'quotation_number', 'contract_number', 'invoice_number',
    'contact_info', 'contact_email', 'contact_phone',
    'company_address', 'created_at', 'updated_at'
}


@https_fn.on_call()
def analyze_template(req: https_fn.CallableRequest) -> dict:
    """
    Analyze a template file to extract variables.

    Request data:
        file_path: str - Storage path to template file

    Returns:
        dict with:
            success: bool
            variables: dict with 'standard' and 'extra' lists
    """

    if not req.auth:
        raise https_fn.HttpsError(
            code=https_fn.FunctionsErrorCode.UNAUTHENTICATED,
            message='Authentication required'
        )

    file_path = req.data.get('file_path')

    if not file_path:
        raise https_fn.HttpsError(
            code=https_fn.FunctionsErrorCode.INVALID_ARGUMENT,
            message='file_path is required'
        )

    try:
        # Download template
        bucket = storage.bucket()
        blob = bucket.blob(file_path)

        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_file:
            blob.download_to_filename(temp_file.name)
            template_file = temp_file.name

        try:
            # Load document and find placeholders
            doc = Document(template_file)
            all_variables = find_placeholders(doc)

            # Categorize into standard and extra
            standard = sorted([v for v in all_variables if v in STANDARD_VARIABLES])
            extra = sorted([v for v in all_variables if v not in STANDARD_VARIABLES])

            return {
                'success': True,
                'variables': {
                    'standard': standard,
                    'extra': extra,
                    'all': sorted(list(all_variables))
                }
            }

        finally:
            import os
            os.unlink(template_file)

    except Exception as e:
        print(f"Error analyzing template: {e}")
        raise https_fn.HttpsError(
            code=https_fn.FunctionsErrorCode.INTERNAL,
            message=f'Internal error: {str(e)}'
        )
